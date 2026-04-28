"""Amazon 用户评论爬虫 - 从用户概况页爬取全部评论及详情"""

import json
import os
import random
import re
import shutil
import sys
import time
import traceback
from datetime import datetime

from playwright.sync_api import sync_playwright
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import queue
import threading
import tkinter as tk
from tkinter import ttk, scrolledtext, messagebox

# ── 路径常量 ──────────────────────────────────────────────────────────────────
BASE_DIR = os.path.dirname(os.path.abspath(sys.argv[0]))

CONFIG_FILE = os.path.join(BASE_DIR, "config.json")
BROWSER_DATA_DIR = os.path.join(BASE_DIR, "browser_data")
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

# ── 默认配置（反爬友好，单 IP 安全） ─────────────────────────────────────────
DEFAULT_CONFIG = {
    "headless": False,
    "page_timeout": 90000,
    "detail_timeout": 30000,
    "delay_min": 4,
    "delay_max": 8,
    "browse_min": 2,
    "browse_max": 5,
    "max_scrolls": 100,
    "max_no_change": 5,
    "user_agent": "",
    "proxy_type": "",
    "proxy_ip": "",
    "proxy_port": "",
}


def load_config():
    """加载配置文件，不存在则返回默认配置"""
    if os.path.exists(CONFIG_FILE):
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            saved = json.load(f)
        merged = {**DEFAULT_CONFIG, **saved}
    else:
        merged = dict(DEFAULT_CONFIG)

    if merged.get("max_no_change", 0) < 1:
        merged["max_no_change"] = 1
    if merged.get("delay_min", 0) < 0:
        merged["delay_min"] = 0
    if merged.get("browse_min", 0) < 0:
        merged["browse_min"] = 0

    return merged


def save_config(cfg):
    """保存配置到文件"""
    with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    print(f"配置已保存到: {CONFIG_FILE}")


def _find_system_chrome_path():
    """查找系统已安装的 Google Chrome 可执行文件路径。"""
    candidates = []

    # 优先检查 Windows 常见安装目录
    if sys.platform.startswith("win"):
        win_bases = [
            os.environ.get("PROGRAMFILES"),
            os.environ.get("PROGRAMFILES(X86)"),
            os.environ.get("LOCALAPPDATA"),
        ]
        for base in win_bases:
            if not base:
                continue
            candidates.append(
                os.path.join(base, "Google", "Chrome", "Application", "chrome.exe")
            )

        # 读取注册表 App Paths，兼容自定义安装路径
        try:
            import winreg

            reg_locations = [
                (winreg.HKEY_CURRENT_USER,
                 r"Software\Microsoft\Windows\CurrentVersion\App Paths\chrome.exe"),
                (winreg.HKEY_LOCAL_MACHINE,
                 r"Software\Microsoft\Windows\CurrentVersion\App Paths\chrome.exe"),
            ]
            for root, subkey in reg_locations:
                try:
                    with winreg.OpenKey(root, subkey) as key:
                        chrome_path, _ = winreg.QueryValueEx(key, None)
                        if chrome_path:
                            candidates.append(chrome_path)
                except OSError:
                    pass
        except Exception:
            pass

    # PATH 中兜底查找
    for cmd in ("chrome.exe", "chrome", "google-chrome", "google-chrome-stable"):
        hit = shutil.which(cmd)
        if hit:
            candidates.append(hit)

    # 按顺序返回第一个存在的路径
    seen = set()
    for path in candidates:
        norm = os.path.normcase(os.path.normpath(path))
        if norm in seen:
            continue
        seen.add(norm)
        if os.path.exists(path):
            return path
    return ""


def _open_persistent_context(pw, cfg, headless=False):
    """使用 user_data_dir 打开持久化浏览器上下文"""
    os.makedirs(BROWSER_DATA_DIR, exist_ok=True)

    for lock_name in ("SingletonLock", "SingletonSocket", "SingletonCookie"):
        lock_path = os.path.join(BROWSER_DATA_DIR, lock_name)
        try:
            if os.path.exists(lock_path):
                os.remove(lock_path)
        except OSError:
            pass

    launch_args = ["--disable-blink-features=AutomationControlled"]
    kwargs = {
        "user_data_dir": BROWSER_DATA_DIR,
        "headless": headless,
        "locale": "en-US",
        "accept_downloads": False,
        "args": launch_args,
    }

    ua = cfg.get("user_agent", "")
    if ua:
        kwargs["user_agent"] = ua

    # 检查是否配置了代理
    proxy_type = cfg.get("proxy_type", "").strip()
    proxy_ip = cfg.get("proxy_ip", "").strip()
    proxy_port = cfg.get("proxy_port", "").strip()

    if proxy_type and proxy_ip and proxy_port:
        proxy_server = f"{proxy_type}://{proxy_ip}:{proxy_port}"
        print(f"✅ 使用代理: {proxy_server}")
        kwargs["proxy"] = {"server": proxy_server}
    else:
        print("ℹ️ 未配置代理，直接访问")

    if headless:
        kwargs["viewport"] = {"width": 1920, "height": 1080}
    else:
        kwargs["no_viewport"] = True
        launch_args.append("--start-maximized")

    chrome_path = _find_system_chrome_path()
    if not chrome_path:
        raise RuntimeError(
            "未检测到系统已安装的 Google Chrome。"
            "请先安装 Chrome 后重试：https://www.google.com/chrome/"
        )

    kwargs["executable_path"] = chrome_path
    print(f"✅ 使用系统 Chrome: {chrome_path}")

    context = pw.chromium.launch_persistent_context(**kwargs)
    return context


def _get_page(context):
    """从 persistent context 获取可用页面"""
    if context.pages:
        page = context.pages[0]
        try:
            page.evaluate("1+1")
            return page
        except Exception:
            pass
    return context.new_page()


# ── Excel 样式常量 ────────────────────────────────────────────────────────────
_THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
_WRAP_ALIGN = Alignment(vertical="top", wrap_text=True)
_CENTER_ALIGN = Alignment(horizontal="center", vertical="top")
_TOP_ALIGN = Alignment(vertical="top")
_STRIPE_FILL = PatternFill(start_color="F2F7FB", end_color="F2F7FB", fill_type="solid")
_GREEN_FONT = Font(color="2E7D32")
_RED_FONT = Font(color="C62828")
_VP_YES_FILL = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
_VP_NO_FILL = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
_RATING_FILLS = {
    5: PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid"),
    4: PatternFill(start_color="F1F8E9", end_color="F1F8E9", fill_type="solid"),
    3: PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid"),
    2: PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid"),
    1: PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid"),
}

# 每列的 (固定宽度, 对齐方式)
_REVIEW_COL_SPEC = {
    "Review ID":          (16, _CENTER_ALIGN),
    "Product Name":       (36, _WRAP_ALIGN),
    "ASIN":               (13, _CENTER_ALIGN),
    "Title":              (30, _WRAP_ALIGN),
    "Description":        (50, _WRAP_ALIGN),
    "Rating":             (8,  _CENTER_ALIGN),
    "Review Date":        (22, _TOP_ALIGN),
    "Verified Purchase":  (12, _CENTER_ALIGN),
    "Helpful Votes":      (10, _CENTER_ALIGN),
    "Hearts":             (8,  _CENTER_ALIGN),
    "Review URL":         (38, _TOP_ALIGN),
}

_REVIEW_HEADERS = list(_REVIEW_COL_SPEC.keys())


def _write_review_sheet(ws, reviews):
    """写入评论数据并应用完整样式"""
    headers = _REVIEW_HEADERS

    # 写表头
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER
    ws.row_dimensions[1].height = 28

    # 写数据行
    for review in reviews:
        ws.append([
            review.get("review_id", ""),
            review.get("product_name", ""),
            review.get("product_asin", ""),
            review.get("title", ""),
            review.get("description", ""),
            review.get("rating", 0),
            review.get("review_date", ""),
            "Yes" if review.get("verified_purchase") else "No",
            review.get("helpful_votes", 0),
            review.get("hearts", 0),
            review.get("review_url", ""),
        ])

    # 样式：逐行处理
    for row_idx in range(2, ws.max_row + 1):
        is_stripe = (row_idx % 2 == 0)
        row_has_wrap = False

        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            spec = _REVIEW_COL_SPEC.get(header)
            if spec:
                cell.alignment = spec[1]
            cell.border = _THIN_BORDER

            # 斑马纹底色（优先级低于条件色）
            if is_stripe:
                cell.fill = _STRIPE_FILL

            # Rating 列条件着色
            if header == "Rating":
                try:
                    rating = int(cell.value) if cell.value else 0
                except (ValueError, TypeError):
                    rating = 0
                if rating in _RATING_FILLS:
                    cell.fill = _RATING_FILLS[rating]
                if rating <= 2:
                    cell.font = Font(bold=True, color="C62828")
                elif rating == 5:
                    cell.font = Font(bold=True, color="2E7D32")

            # VP 列条件着色
            if header == "Verified Purchase":
                if cell.value == "Yes":
                    cell.fill = _VP_YES_FILL
                    cell.font = _GREEN_FONT
                elif cell.value == "No":
                    cell.fill = _VP_NO_FILL
                    cell.font = _RED_FONT

            # 检查是否需要增大行高
            if spec and spec[1].wrap_text and cell.value:
                text = str(cell.value)
                col_w = spec[0]
                chars_per_line = int(col_w * 1.2)
                if chars_per_line > 0 and len(text) > chars_per_line:
                    row_has_wrap = True

        # 设置行高
        if row_has_wrap:
            ws.row_dimensions[row_idx].height = 45
        else:
            ws.row_dimensions[row_idx].height = 22

    # 列宽
    for col_idx, header in enumerate(headers, 1):
        letter = get_column_letter(col_idx)
        spec = _REVIEW_COL_SPEC.get(header)
        ws.column_dimensions[letter].width = spec[0] if spec else 15

    # 冻结首行 + 自动筛选
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _build_analysis_sheet(wb, all_data):
    """构建分析报告 sheet，带完整样式"""
    analysis = _analyze_users(all_data)
    ws = wb.create_sheet(title="分析报告")

    a_headers = [
        "用户名", "评论总数", "5星占比", "VP占比",
        "平均评论长度", "有帮助投票", "可疑评分", "结论", "详细原因",
    ]
    a_widths = [18, 10, 10, 10, 14, 12, 10, 18, 55]

    # 表头
    ws.append(a_headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGN
        cell.border = _THIN_BORDER
    ws.row_dimensions[1].height = 28

    # 评分阈值颜色
    score_fills = {
        "high": PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid"),
        "mid": PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid"),
        "low": PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid"),
    }

    for row_i, r in enumerate(analysis, 2):
        ws.append([
            r["name"],
            r["total"],
            f"{r['five_pct']:.0f}%",
            f"{r['vp_pct']:.0f}%",
            f"{r['avg_len']:.0f}",
            r["total_helpful"],
            r["score"],
            r["conclusion"],
            "; ".join(r["reasons"]) if r["reasons"] else "无异常",
        ])

        is_stripe = (row_i % 2 == 0)
        for col_idx in range(1, len(a_headers) + 1):
            cell = ws.cell(row=row_i, column=col_idx)
            cell.border = _THIN_BORDER
            cell.alignment = _CENTER_ALIGN if col_idx <= 7 else _WRAP_ALIGN
            if is_stripe:
                cell.fill = _STRIPE_FILL

        # 可疑评分 + 结论列着色
        score_cell = ws.cell(row=row_i, column=7)
        conclusion_cell = ws.cell(row=row_i, column=8)
        try:
            sc = int(score_cell.value) if score_cell.value else 0
        except (ValueError, TypeError):
            sc = 0
        if sc >= 5:
            score_cell.fill = score_fills["high"]
            score_cell.font = Font(bold=True, color="C62828")
            conclusion_cell.fill = score_fills["high"]
            conclusion_cell.font = Font(bold=True, color="C62828")
        elif sc >= 3:
            score_cell.fill = score_fills["mid"]
            score_cell.font = Font(bold=True, color="F57F17")
            conclusion_cell.fill = score_fills["mid"]
            conclusion_cell.font = Font(bold=True, color="F57F17")
        else:
            score_cell.fill = score_fills["low"]
            score_cell.font = Font(bold=True, color="2E7D32")
            conclusion_cell.fill = score_fills["low"]
            conclusion_cell.font = Font(bold=True, color="2E7D32")

        # 详细原因列自动行高
        reason_text = str(ws.cell(row=row_i, column=9).value or "")
        if len(reason_text) > 50:
            ws.row_dimensions[row_i].height = 38
        else:
            ws.row_dimensions[row_i].height = 22

    # 规则说明区
    gap_row = len(analysis) + 3
    ws.cell(row=gap_row, column=1).value = "评分规则说明:"
    ws.cell(row=gap_row, column=1).font = Font(bold=True, size=11, color="333333")
    ws.row_dimensions[gap_row].height = 26

    rules = [
        ("5星评价占比 > 80%", "+2分", "高比例好评可能为虚假评论"),
        ("已验证购买占比 < 40%", "+2分", "大量评论无购买记录"),
        ("评论平均长度 < 80字符", "+1分", "内容空洞，可能为模板评论"),
        ("同日发布 ≥ 3条评论", "+2分", "短时间大量发布"),
        ("所有评论无人点赞(≥5条时)", "+1分", "评论缺乏互动"),
        ("评论总数 > 50条", "+1分", "活跃度异常高"),
    ]

    rule_header_row = gap_row + 1
    for col_idx, h in enumerate(["规则条件", "分值", "说明"], 1):
        cell = ws.cell(row=rule_header_row, column=col_idx)
        cell.value = h
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill(start_color="7B8FA1", end_color="7B8FA1", fill_type="solid")
        cell.alignment = _CENTER_ALIGN
        cell.border = _THIN_BORDER
    ws.row_dimensions[rule_header_row].height = 22

    for ri, (cond, score_txt, desc) in enumerate(rules, rule_header_row + 1):
        ws.cell(row=ri, column=1, value=cond)
        ws.cell(row=ri, column=2, value=score_txt)
        ws.cell(row=ri, column=3, value=desc)
        for ci in range(1, 4):
            cell = ws.cell(row=ri, column=ci)
            cell.border = _THIN_BORDER
            cell.alignment = _CENTER_ALIGN if ci == 2 else _TOP_ALIGN
            if ri % 2 == 0:
                cell.fill = _STRIPE_FILL
        ws.row_dimensions[ri].height = 20

    # 结论说明
    legend_row = rule_header_row + len(rules) + 1
    ws.cell(row=legend_row, column=1, value="结论判定:").font = Font(bold=True, size=10)
    ws.cell(row=legend_row + 1, column=1, value="0-2分 → ✅ 正常")
    ws.cell(row=legend_row + 1, column=1).font = Font(color="2E7D32")
    ws.cell(row=legend_row + 2, column=1, value="3-4分 → ⚠️ 存在嫌疑")
    ws.cell(row=legend_row + 2, column=1).font = Font(color="F57F17")
    ws.cell(row=legend_row + 3, column=1, value="≥5分  → 🚨 高度疑似")
    ws.cell(row=legend_row + 3, column=1).font = Font(color="C62828")

    # 列宽
    for col_idx, w in enumerate(a_widths, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # 冻结首行 + 自动筛选
    ws.freeze_panes = "A2"
    filter_end = get_column_letter(len(a_headers))
    ws.auto_filter.ref = f"A1:{filter_end}{len(analysis) + 1}"


class AmazonReviewScraper:
    def __init__(self, profile_url, config=None):
        self.profile_url = profile_url
        self.cfg = config or load_config()

    def _random_delay(self, low=None, high=None):
        low = low if low is not None else self.cfg["delay_min"]
        high = high if high is not None else self.cfg["delay_max"]
        time.sleep(random.uniform(low, high))

    def _check_captcha(self, page):
        captcha_indicators = [
            "captcha", "robot", "automated access",
            "Type the characters you see in this image",
        ]
        body_text = page.evaluate("() => document.body.innerText || ''")
        for indicator in captcha_indicators:
            if indicator.lower() in body_text.lower():
                print("\n⚠️  检测到 Amazon 人机验证（验证码）！")
                cb = getattr(self, '_captcha_callback', None)
                if cb:
                    cb()
                else:
                    print("请在浏览器窗口中手动完成验证，然后按回车继续...")
                    input()
                return True
        return False

    def _load_profile_page(self, page):
        print(f"正在打开概况页: {self.profile_url}")
        page.goto(self.profile_url, wait_until="domcontentloaded",
                  timeout=self.cfg["page_timeout"])
        page.wait_for_timeout(3000)
        self._check_captcha(page)

        max_scrolls = self.cfg["max_scrolls"]
        max_no_change = self.cfg["max_no_change"]
        no_change_count = 0
        prev_count = 0

        for i in range(max_scrolls):
            current_count = page.evaluate("""
                () => {
                    const script = document.querySelector('script[data-a-state*="page-state-profile"]');
                    if (!script) return 0;
                    try {
                        const data = JSON.parse(script.textContent);
                        const items = data.reviewsTimeline && data.reviewsTimeline.shopItemModels;
                        return items ? items.length : 0;
                    } catch(e) { return 0; }
                }
            """)

            print(f"  滚动第 {i + 1} 次，当前评论数: {current_count}")

            if current_count == prev_count:
                no_change_count += 1
                if no_change_count >= max_no_change:
                    print(f"  连续 {max_no_change} 次无新评论，停止滚动")
                    break
            else:
                no_change_count = 0

            prev_count = current_count
            page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
            self._random_delay(2, 3)

        print(f"评论加载完毕，共 {prev_count} 条")

    def _extract_user_info(self, page):
        user_info = {"name": "", "entity_id": "", "avatar": ""}

        try:
            name_el = page.query_selector("#shop-influencer-profile-name")
            if name_el:
                user_info["name"] = name_el.inner_text().strip()
        except Exception:
            pass

        try:
            avatar_el = page.query_selector("#shop-influencer-profile-image")
            if avatar_el:
                user_info["avatar"] = avatar_el.get_attribute("src") or ""
        except Exception:
            pass

        match = re.search(r"amzn1\.account\.[A-Z0-9]+", self.profile_url)
        if match:
            user_info["entity_id"] = match.group(0)

        if not user_info["entity_id"]:
            try:
                entity_id = page.evaluate("""
                    () => {
                        const scripts = document.querySelectorAll('script');
                        for (const s of scripts) {
                            const match = s.textContent.match(/"directedId"\\s*:\\s*"([^"]+)"/);
                            if (match) return match[1];
                        }
                        return '';
                    }
                """)
                user_info["entity_id"] = entity_id
            except Exception:
                pass

        print(f"用户信息: {user_info['name']} ({user_info['entity_id']})")
        return user_info

    def _extract_reviews_from_json(self, page):
        script_content = page.evaluate("""
            () => {
                const script = document.querySelector('script[data-a-state*="page-state-profile"]');
                return script ? script.textContent : null;
            }
        """)

        if not script_content:
            raise RuntimeError("无法在概况页中找到评论数据的 JSON 脚本标签")

        data = json.loads(script_content)
        shop_items = data.get("reviewsTimeline", {}).get("shopItemModels", [])

        if not shop_items:
            print("警告: 未在概况页 JSON 中找到任何评论")
            return []

        reviews = []
        for item in shop_items:
            rm = item.get("reviewModel", {})
            if not rm:
                continue

            view_url = rm.get("viewOnAmazon", "")
            review_id = self._extract_review_id(view_url)

            review = {
                "title": rm.get("title", ""),
                "description": rm.get("description", ""),
                "rating": rm.get("rating", 0),
                "product_name": rm.get("asinTitle", ""),
                "product_asin": self._extract_asin(view_url),
                "product_image": rm.get("productImageUrl", ""),
                "review_url": view_url.split("?")[0] if view_url else "",
                "review_id": review_id,
                "hearts": rm.get("hearts", 0),
                "helpful_votes": self._parse_helpful_votes(
                    rm.get("helpfulVoteText", "")
                ),
                "review_date": "",
                "verified_purchase": False,
            }
            reviews.append(review)

        print(f"从概况页 JSON 中解析到 {len(reviews)} 条评论")
        return reviews

    @staticmethod
    def _extract_review_id(view_on_amazon_url):
        if not view_on_amazon_url:
            return ""
        match = re.search(r"/customer-reviews/([A-Z0-9]+)", view_on_amazon_url)
        return match.group(1) if match else ""

    @staticmethod
    def _extract_asin(url):
        if not url:
            return ""
        match = re.search(r"/dp/([A-Z0-9]{10})", url)
        return match.group(1) if match else ""

    @staticmethod
    def _parse_helpful_votes(text):
        if not text:
            return 0
        match = re.search(r"(\d+)", text)
        return int(match.group(1)) if match else 0

    def _simulate_browsing(self, page):
        browse_min = self.cfg.get("browse_min", 2)
        browse_max = self.cfg.get("browse_max", 5)
        page.wait_for_timeout(800)
        scroll_count = random.randint(2, 4)
        for _ in range(scroll_count):
            scroll_amount = random.randint(200, 500)
            page.evaluate(f"window.scrollBy(0, {scroll_amount})")
            time.sleep(random.uniform(0.5, 1.5))
        time.sleep(random.uniform(browse_min, browse_max))

    def _scrape_detail_in_new_tab(self, context, review_url):
        result = {"review_date": "获取失败", "verified_purchase": False}
        if not review_url:
            return result

        detail_page = None
        try:
            detail_page = context.new_page()
            detail_page.goto(review_url, wait_until="domcontentloaded",
                             timeout=self.cfg["detail_timeout"])
            self._check_captcha(detail_page)
            self._simulate_browsing(detail_page)

            try:
                date_el = detail_page.query_selector('[data-hook="review-date"]')
                if date_el:
                    result["review_date"] = date_el.inner_text().strip()
            except Exception:
                pass

            try:
                avp_el = detail_page.query_selector('[data-hook="avp-badge"]')
                result["verified_purchase"] = avp_el is not None
            except Exception:
                pass

        except Exception as e:
            print(f"    访问详情页失败: {e}")
        finally:
            if detail_page:
                try:
                    detail_page.close()
                except Exception:
                    pass

        return result

    def scrape(self, progress_callback=None, captcha_callback=None):
        self._captcha_callback = captcha_callback
        with sync_playwright() as p:
            headless = self.cfg["headless"]
            has_data = os.path.exists(BROWSER_DATA_DIR)
            print(f"正在启动浏览器（{'无头' if headless else '有头'}模式）...")
            if has_data:
                print(f"已加载浏览器数据: {BROWSER_DATA_DIR}")
            else:
                print("未找到浏览器数据，将以未登录身份爬取")

            context = _open_persistent_context(p, self.cfg, headless=headless)
            profile_page = _get_page(context)

            try:
                self._load_profile_page(profile_page)
                user_info = self._extract_user_info(profile_page)
                reviews = self._extract_reviews_from_json(profile_page)

                total = len(reviews)
                start_time = time.time()

                for idx, review in enumerate(reviews, 1):
                    review_url = review["review_url"]
                    elapsed = time.time() - start_time
                    if idx > 1:
                        avg_per_review = elapsed / (idx - 1)
                        remaining = avg_per_review * (total - idx + 1)
                        eta_min = int(remaining // 60)
                        eta_sec = int(remaining % 60)
                        eta_str = f"  预计剩余 {eta_min}分{eta_sec}秒"
                    else:
                        eta_str = ""

                    print(f"[{idx}/{total}] 正在获取详情: "
                          f"{review['review_id']}{eta_str}")

                    if progress_callback:
                        progress_callback(idx, total, elapsed)

                    if idx > 1:
                        self._random_delay()

                    detail = self._scrape_detail_in_new_tab(context, review_url)
                    review["review_date"] = detail["review_date"]
                    review["verified_purchase"] = detail["verified_purchase"]
                    print(f"    日期: {review['review_date']}, "
                          f"VP: {review['verified_purchase']}")

                total_time = time.time() - start_time
                print(f"\n爬取耗时: {int(total_time // 60)}分"
                      f"{int(total_time % 60)}秒")

                result = {
                    "user": user_info,
                    "scraped_at": datetime.now().isoformat(timespec="seconds"),
                    "total_reviews": len(reviews),
                    "reviews": reviews,
                }
                return result

            finally:
                context.close()
                print("浏览器已关闭")

    @staticmethod
    def save_to_json(data, output_dir=None):
        output_dir = output_dir or OUTPUT_DIR
        os.makedirs(output_dir, exist_ok=True)

        username = data.get("user", {}).get("name", "unknown") or "unknown"
        safe_name = re.sub(r'[\\/:*?"<>|]', "_", username)
        filename = f"{safe_name}_reviews.json"
        filepath = os.path.join(output_dir, filename)

        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        print(f"数据已保存到: {filepath}")
        return filepath

    @staticmethod
    def save_to_excel(data, output_dir=None):
        output_dir = output_dir or OUTPUT_DIR
        os.makedirs(output_dir, exist_ok=True)

        username = data.get("user", {}).get("name", "unknown") or "unknown"
        safe_name = re.sub(r'[\\/:*?"<>|]', "_", username)
        filename = f"{safe_name}_reviews.xlsx"
        filepath = os.path.join(output_dir, filename)

        wb = Workbook()
        ws = wb.active
        ws.title = "Reviews"
        _write_review_sheet(ws, data.get("reviews", []))

        wb.save(filepath)
        print(f"Excel 已保存到: {filepath}")
        return filepath

    @staticmethod
    def save_batch_excel(all_data, output_dir=None):
        """将多个用户的数据保存为一个 Excel 文件，含分析报告"""
        output_dir = output_dir or OUTPUT_DIR
        os.makedirs(output_dir, exist_ok=True)

        ts = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"batch_{ts}.xlsx"
        filepath = os.path.join(output_dir, filename)

        wb = Workbook()

        for i, data in enumerate(all_data):
            user = data.get("user", {})
            name = user.get("name", "未知") or "未知"
            safe = re.sub(r'[\\/:*?"<>|\[\]]', "_", name)[:28]
            sheet_name = safe if i == 0 else f"{safe}_{i}"

            ws = wb.active if i == 0 else wb.create_sheet()
            ws.title = sheet_name
            _write_review_sheet(ws, data.get("reviews", []))

        # ── 分析报告 sheet ──
        _build_analysis_sheet(wb, all_data)

        wb.save(filepath)
        print(f"批量 Excel 已保存到: {filepath}")
        return filepath


def _analyze_users(all_data):
    """分析多个用户的评论模式，检测刷评嫌疑"""
    results = []
    for data in all_data:
        user = data.get("user", {})
        reviews = data.get("reviews", [])
        name = user.get("name", "未知") or "未知"
        total = len(reviews)

        if total == 0:
            results.append({
                "name": name, "total": 0, "five_pct": 0, "vp_pct": 0,
                "avg_len": 0, "total_helpful": 0, "score": 0,
                "conclusion": "无评论数据", "reasons": [],
            })
            continue

        reasons = []
        score = 0

        five_star = sum(1 for r in reviews if r.get("rating", 0) == 5)
        five_pct = five_star / total * 100
        if five_pct > 80 and total >= 5:
            score += 2
            reasons.append(f"5星评价占比 {five_pct:.0f}%，远高于正常水平")

        vp_count = sum(1 for r in reviews if r.get("verified_purchase"))
        vp_pct = vp_count / total * 100
        if vp_pct < 40:
            score += 2
            reasons.append(f"已验证购买仅占 {vp_pct:.0f}%，大量评论无购买记录")

        desc_lens = [len(r.get("description", "")) for r in reviews]
        avg_len = sum(desc_lens) / total
        if avg_len < 80:
            score += 1
            reasons.append(f"评论平均长度仅 {avg_len:.0f} 字符，内容空洞")

        date_counts = {}
        for r in reviews:
            rd = r.get("review_date", "")
            m = re.search(r"on (\w+ \d+, \d{4})", rd)
            if m:
                d = m.group(1)
                date_counts[d] = date_counts.get(d, 0) + 1
        burst = {d: c for d, c in date_counts.items() if c >= 3}
        if burst:
            top = sorted(burst.items(), key=lambda x: -x[1])[:3]
            info = ", ".join(f"{d}({c}条)" for d, c in top)
            score += 2
            reasons.append(f"同日大量评论: {info}")

        total_helpful = sum(r.get("helpful_votes", 0) for r in reviews)
        if total_helpful == 0 and total >= 5:
            score += 1
            reasons.append("所有评论均无人认为有帮助")

        if total > 50:
            score += 1
            reasons.append(f"评论数量达 {total} 条，活跃度异常高")

        if score >= 5:
            conclusion = "🚨 高度疑似刷评"
        elif score >= 3:
            conclusion = "⚠️ 存在刷评嫌疑"
        else:
            conclusion = "✅ 未发现明显异常"

        results.append({
            "name": name, "total": total,
            "five_pct": five_pct, "vp_pct": vp_pct,
            "avg_len": avg_len, "total_helpful": total_helpful,
            "score": score, "conclusion": conclusion,
            "reasons": reasons,
        })
    return results


# ── 登录模式 ──────────────────────────────────────────────────────────────────
def do_login():
    cfg = load_config()
    print("=" * 50)
    print("  Amazon 登录模式")
    print("=" * 50)
    print()
    print("即将打开浏览器窗口，请手动登录您的 Amazon 账号。")
    print(f"浏览器数据保存位置: {BROWSER_DATA_DIR}")
    print()
    print("登录完成后，回到此窗口按回车关闭浏览器并保存状态。")
    print()

    with sync_playwright() as p:
        context = _open_persistent_context(p, cfg, headless=False)
        page = _get_page(context)
        page.goto(
            "https://www.amazon.com/ap/signin"
            "?openid.pape.max_auth_age=0"
            "&openid.return_to=https%3A%2F%2Fwww.amazon.com%2F"
            "&openid.identity=http%3A%2F%2Fspecs.openid.net%2Fauth%2F2.0%2Fidentifier_select"
            "&openid.assoc_handle=usflex"
            "&openid.mode=checkid_setup"
            "&openid.claimed_id=http%3A%2F%2Fspecs.openid.net%2Fauth%2F2.0%2Fidentifier_select"
            "&openid.ns=http%3A%2F%2Fspecs.openid.net%2Fauth%2F2.0",
            timeout=cfg["page_timeout"],
        )
        input("登录完成后，请按回车保存状态并关闭浏览器...")
        context.close()

    print()
    print(f"✅ 浏览器状态已保存到: {BROWSER_DATA_DIR}")
    print("下次爬取时将自动使用此登录状态。")


# ── 爬取模式 ──────────────────────────────────────────────────────────────────
def do_scrape():
    cfg = load_config()
    print("=" * 50)
    print("  Amazon 评论爬取模式")
    print("=" * 50)

    choice = input(f"是否使用无头模式？(y/n, 当前默认: "
                   f"{'无头' if cfg['headless'] else '有头'}): ").strip().lower()
    if choice == "y":
        cfg["headless"] = True
    elif choice == "n":
        cfg["headless"] = False

    print("请输入 Amazon 用户概况页 URL（每行一个，输入空行结束）:")
    urls = []
    while True:
        line = input("  URL: ").strip()
        if not line:
            break
        urls.append(line)

    if not urls:
        print("未输入任何 URL")
        return

    all_data = []
    for i, url in enumerate(urls, 1):
        print(f"\n══ 用户 {i}/{len(urls)} ══")
        if "amazon.com/gp/profile/" not in url:
            print("警告: URL 看起来不是 Amazon 用户概况页，仍将尝试...")
        scraper = AmazonReviewScraper(profile_url=url, config=cfg)
        try:
            data = scraper.scrape()
            AmazonReviewScraper.save_to_json(data)
            all_data.append(data)
        except Exception as e:
            print(f"\n用户 {i} 爬取失败: {e}")
            traceback.print_exc()

    if all_data:
        xlsx_path = AmazonReviewScraper.save_batch_excel(all_data)
        print(f"\n爬取完成！成功 {len(all_data)}/{len(urls)} 个用户")
        print(f"Excel 报告: {xlsx_path}")
    else:
        print("\n所有用户爬取均失败")


# ── 配置模式 ──────────────────────────────────────────────────────────────────
def do_config():
    cfg = load_config()
    print("=" * 50)
    print("  配置管理")
    print("=" * 50)
    print("直接回车保持当前值不变。\n")

    def ask(prompt, current, cast=str):
        val = input(f"  {prompt} [{current}]: ").strip()
        if not val:
            return current
        try:
            return cast(val)
        except ValueError:
            print(f"    输入无效，保持原值: {current}")
            return current

    def ask_bool(prompt, current):
        val = input(f"  {prompt} (y/n) [{'y' if current else 'n'}]: ").strip().lower()
        if val == "y":
            return True
        elif val == "n":
            return False
        return current

    cfg["headless"] = ask_bool("默认无头模式", cfg["headless"])
    cfg["page_timeout"] = ask("概况页超时 (毫秒)", cfg["page_timeout"], int)
    cfg["detail_timeout"] = ask("详情页超时 (毫秒)", cfg["detail_timeout"], int)
    cfg["delay_min"] = ask("评论间隔最小值 (秒)", cfg["delay_min"], float)
    cfg["delay_max"] = ask("评论间隔最大值 (秒)", cfg["delay_max"], float)
    cfg["browse_min"] = ask("详情页浏览最小时间 (秒)", cfg["browse_min"], float)
    cfg["browse_max"] = ask("详情页浏览最大时间 (秒)", cfg["browse_max"], float)
    cfg["max_scrolls"] = ask("最大滚动次数", cfg["max_scrolls"], int)
    cfg["max_no_change"] = ask("无新数据最大重试次数", cfg["max_no_change"], int)
    cfg["user_agent"] = ask("代理协议",cfg["proxy_type"] or "")
    cfg["user_agent"] = ask("代理 IP",cfg["proxy_ip"] or "")
    cfg["user_agent"] = ask("代理端口",cfg["proxy_port"] or "")
    cfg["user_agent"] = ask("User-Agent (留空=使用浏览器默认)",
                            cfg["user_agent"] or "(浏览器默认)")
    if cfg["user_agent"] == "(浏览器默认)":
        cfg["user_agent"] = ""
    save_config(cfg)


def do_reset_config():
    save_config(dict(DEFAULT_CONFIG))
    print("已重置为默认配置")


# ── 主菜单 ────────────────────────────────────────────────────────────────────
def show_menu():
    print()
    print("=" * 50)
    print("  Amazon 用户评论爬虫")
    print("=" * 50)
    print()
    print("  1. 登录 Amazon（保存浏览器状态）")
    print("  2. 爬取用户评论")
    print("  3. 修改配置")
    print("  4. 重置为默认配置")
    print("  0. 退出")
    print()
    has_data = os.path.exists(BROWSER_DATA_DIR)
    has_config = os.path.exists(CONFIG_FILE)
    print(f"  登录状态: {'✅ 已保存' if has_data else '❌ 未登录'}")
    print(f"  配置文件: {'✅ 已存在' if has_config else '使用默认配置'}")
    print()
    return input("请选择操作 [1/2/3/4/0]: ").strip()


# ── GUI 界面 ──────────────────────────────────────────────────────────────────


class _QueueWriter:
    def __init__(self, q):
        self._q = q

    def write(self, text):
        if text:
            self._q.put(("log", text))

    def flush(self):
        pass


class SettingsDialog(tk.Toplevel):
    def __init__(self, parent):
        super().__init__(parent)
        self.title("设置")
        self.resizable(False, False)
        self.grab_set()

        cfg = load_config()
        self._entries = {}

        frame = ttk.Frame(self, padding=15)
        frame.pack(fill=tk.BOTH, expand=True)

        fields = [
            ("headless", "默认无头模式", "bool"),
            ("page_timeout", "概况页超时 (毫秒)", "int"),
            ("detail_timeout", "详情页超时 (毫秒)", "int"),
            ("delay_min", "评论间隔最小值 (秒)", "float"),
            ("delay_max", "评论间隔最大值 (秒)", "float"),
            ("browse_min", "详情页浏览最小时间 (秒)", "float"),
            ("browse_max", "详情页浏览最大时间 (秒)", "float"),
            ("max_scrolls", "最大滚动次数", "int"),
            ("max_no_change", "无新数据最大重试次数", "int"),
            ("user_agent", "User-Agent (留空=浏览器默认)", "str"),
            # ("proxy_type", "代理协议", "str"),
            ("proxy_ip", "代理 IP", "str"),
            ("proxy_port", "代理端口", "str"),
        ]

        for row, (key, label, typ) in enumerate(fields):
            ttk.Label(frame, text=label).grid(
                row=row, column=0, sticky=tk.W, pady=2)
            if typ == "bool":
                var = tk.BooleanVar(value=cfg.get(key, False))
                ttk.Checkbutton(frame, variable=var).grid(
                    row=row, column=1, sticky=tk.W, pady=2)
                self._entries[key] = ("bool", var)
            else:
                var = tk.StringVar(value=str(cfg.get(key, "")))
                ttk.Entry(frame, textvariable=var, width=30).grid(
                    row=row, column=1, sticky=tk.EW, pady=2, padx=(10, 0))
                self._entries[key] = (typ, var)

        proxy_row = len(fields)
        ttk.Label(frame, text="代理协议").grid(
            row=proxy_row, column=0, sticky=tk.W, pady=2)
        proxy_var = tk.StringVar(value=cfg.get("proxy_type", ""))
        proxy_combo = ttk.Combobox(
            frame, textvariable=proxy_var, width=28, state="readonly")
        proxy_combo["values"] = ("", "http", "https", "socks4", "socks5")
        proxy_combo.grid(row=proxy_row, column=1, sticky=tk.EW, pady=2, padx=(10, 0))
        self._entries["proxy_type"] = ("str", proxy_var)


        btn_frame = ttk.Frame(frame)
        btn_frame.grid(row=len(fields)+1, column=0, columnspan=2, pady=(15, 0))
        ttk.Button(btn_frame, text="保存",
                   command=self._save).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="取消",
                   command=self.destroy).pack(side=tk.LEFT, padx=5)

        self.transient(parent)
        self.update_idletasks()
        x = parent.winfo_x() + (parent.winfo_width() - self.winfo_width()) // 2
        y = parent.winfo_y() + (parent.winfo_height() - self.winfo_height()) // 2
        self.geometry(f"+{x}+{y}")

    def _save(self):
        cfg = {}
        for key, (typ, var) in self._entries.items():
            try:
                if typ == "bool":
                    cfg[key] = var.get()
                elif typ == "int":
                    cfg[key] = int(var.get())
                elif typ == "float":
                    cfg[key] = float(var.get())
                else:
                    cfg[key] = var.get()
            except (ValueError, tk.TclError):
                messagebox.showerror(
                    "输入错误", f"字段 '{key}' 的值无效", parent=self)
                return
        save_config(cfg)
        messagebox.showinfo("设置", "配置已保存", parent=self)
        self.destroy()


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Amazon 用户评论爬虫")
        self.geometry("720x620")
        self.minsize(620, 540)

        self._msg_queue = queue.Queue()
        self._working = False
        self._batch_prefix = ""
        self._destroyed = False

        self._build_ui()
        self._redirect_stdout()
        self._poll_queue()
        self._update_login_status()
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def _build_ui(self):
        pad = {"padx": 10, "pady": 4}

        ttk.Label(self, text="Amazon 用户评论爬虫",
                  font=("", 16, "bold")).pack(pady=(12, 4))

        login_frm = ttk.LabelFrame(self, text="登录", padding=8)
        login_frm.pack(fill=tk.X, **pad)
        self._login_status = ttk.Label(login_frm, text="")
        self._login_status.pack(side=tk.LEFT, padx=(0, 10))
        self._login_btn = ttk.Button(
            login_frm, text="登录 Amazon", command=self._on_login)
        self._login_btn.pack(side=tk.RIGHT)

        scrape_frm = ttk.LabelFrame(self, text="爬取", padding=8)
        scrape_frm.pack(fill=tk.X, **pad)

        ttk.Label(scrape_frm, text="概况页 URL（每行一个，支持批量）:").pack(
            anchor=tk.W)
        self._url_text = tk.Text(scrape_frm, height=4, font=("Consolas", 9))
        self._url_text.pack(fill=tk.X, pady=(4, 0))

        ctrl_row = ttk.Frame(scrape_frm)
        ctrl_row.pack(fill=tk.X, pady=(6, 0))
        self._headless_var = tk.BooleanVar(
            value=load_config().get("headless", False))
        ttk.Checkbutton(ctrl_row, text="无头模式",
                        variable=self._headless_var).pack(side=tk.LEFT)
        self._scrape_btn = ttk.Button(
            ctrl_row, text="开始爬取", command=self._on_scrape)
        self._scrape_btn.pack(side=tk.RIGHT)

        prog_frm = ttk.Frame(self)
        prog_frm.pack(fill=tk.X, **pad)
        self._progress = ttk.Progressbar(prog_frm, mode="determinate")
        self._progress.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self._prog_label = ttk.Label(
            prog_frm, text="就绪", width=28, anchor=tk.E)
        self._prog_label.pack(side=tk.RIGHT, padx=(8, 0))

        self._log = scrolledtext.ScrolledText(
            self, height=14, state=tk.DISABLED,
            font=("Consolas", 9), wrap=tk.WORD)
        self._log.pack(fill=tk.BOTH, expand=True, **pad)

        bottom = ttk.Frame(self)
        bottom.pack(fill=tk.X, padx=10, pady=(0, 10))
        ttk.Button(bottom, text="⚙ 设置",
                   command=self._on_settings).pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(bottom, text="重置配置",
                   command=self._on_reset).pack(side=tk.LEFT)

    def _redirect_stdout(self):
        self._orig_stdout = sys.stdout
        sys.stdout = _QueueWriter(self._msg_queue)

    def _poll_queue(self):
        if self._destroyed:
            return

        while True:
            try:
                msg = self._msg_queue.get_nowait()
            except queue.Empty:
                break

            if not isinstance(msg, tuple):
                self._append_log(str(msg))
                continue

            kind = msg[0]
            if kind == "log":
                self._append_log(msg[1])
            elif kind == "user_progress":
                ui, ut = msg[1], msg[2]
                self._batch_prefix = f"用户{ui}/{ut} | " if ut > 1 else ""
                self._progress["value"] = 0
                self._prog_label.config(
                    text=f"{self._batch_prefix}准备中...")
            elif kind == "progress":
                idx, total, elapsed = msg[1], msg[2], msg[3]
                self._progress["maximum"] = total
                self._progress["value"] = idx
                pfx = self._batch_prefix
                if idx > 1 and elapsed > 0:
                    avg = elapsed / (idx - 1)
                    rem = avg * (total - idx)
                    m, s = int(rem // 60), int(rem % 60)
                    self._prog_label.config(
                        text=f"{pfx}{idx}/{total} 约剩{m}分{s}秒")
                else:
                    self._prog_label.config(text=f"{pfx}{idx}/{total}")
            elif kind == "captcha":
                event = msg[1]
                messagebox.showinfo(
                    "验证码",
                    "检测到 Amazon 人机验证！\n"
                    "请在浏览器中完成验证，然后点击确定继续。",
                    parent=self)
                event.set()
            elif kind == "login_wait":
                event = msg[1]
                messagebox.showinfo(
                    "登录",
                    "请在浏览器窗口中登录 Amazon。\n"
                    "登录完成后点击确定保存状态。",
                    parent=self)
                event.set()
            elif kind == "done":
                self._set_working(False)
                self._update_login_status()
                if len(msg) > 1 and msg[1]:
                    messagebox.showinfo("完成", msg[1], parent=self)

        if not self._destroyed:
            self.after(100, self._poll_queue)

    def _append_log(self, text):
        self._log.config(state=tk.NORMAL)
        self._log.insert(tk.END, text)
        self._log.see(tk.END)
        self._log.config(state=tk.DISABLED)

    def _set_working(self, working):
        self._working = working
        state = "disabled" if working else "normal"
        self._login_btn.config(state=state)
        self._scrape_btn.config(state=state)
        if not working:
            self._batch_prefix = ""
            self._progress["value"] = 0
            self._prog_label.config(text="就绪")

    def _update_login_status(self):
        has = os.path.exists(BROWSER_DATA_DIR)
        self._login_status.config(
            text="✅ 已登录" if has else "❌ 未登录")

    def _on_close(self):
        if self._working:
            if not messagebox.askyesno(
                    "确认", "任务正在运行中，确定要退出吗？", parent=self):
                return
        self.destroy()

    def destroy(self):
        self._destroyed = True
        if hasattr(self, '_orig_stdout'):
            sys.stdout = self._orig_stdout
        super().destroy()

    # ── 登录 ──

    def _on_login(self):
        self._set_working(True)
        self._append_log("\n── 开始登录 ──\n")
        threading.Thread(target=self._login_worker, daemon=True).start()

    def _login_worker(self):
        try:
            cfg = load_config()
            print("正在启动浏览器...\n")
            with sync_playwright() as p:
                context = _open_persistent_context(p, cfg, headless=False)
                page = _get_page(context)
                page.goto(
                    "https://www.amazon.com/ap/signin"
                    "?openid.pape.max_auth_age=0"
                    "&openid.return_to=https%3A%2F%2Fwww.amazon.com%2F"
                    "&openid.identity=http%3A%2F%2Fspecs.openid.net"
                    "%2Fauth%2F2.0%2Fidentifier_select"
                    "&openid.assoc_handle=usflex"
                    "&openid.mode=checkid_setup"
                    "&openid.claimed_id=http%3A%2F%2Fspecs.openid.net"
                    "%2Fauth%2F2.0%2Fidentifier_select"
                    "&openid.ns=http%3A%2F%2Fspecs.openid.net"
                    "%2Fauth%2F2.0",
                    timeout=cfg["page_timeout"],
                )
                wait_event = threading.Event()
                self._msg_queue.put(("login_wait", wait_event))
                wait_event.wait()
                context.close()
            print("✅ 浏览器状态已保存\n")
            self._msg_queue.put((
                "done",
                "登录状态已保存！\n下次爬取将自动使用此登录。"))
        except Exception as e:
            print(f"登录失败: {e}\n")
            traceback.print_exc()
            self._msg_queue.put(("done",))

    # ── 爬取 ──

    def _on_scrape(self):
        raw = self._url_text.get("1.0", tk.END).strip()
        urls = [u.strip() for u in raw.splitlines() if u.strip()]
        if not urls:
            messagebox.showwarning(
                "提示", "请输入至少一个 Amazon 用户概况页 URL", parent=self)
            return
        headless = self._headless_var.get()
        self._set_working(True)
        self._append_log(f"\n── 开始爬取（{len(urls)} 个用户）──\n")
        threading.Thread(
            target=self._scrape_worker,
            args=(urls, headless), daemon=True).start()

    def _scrape_worker(self, urls, headless):
        try:
            cfg = load_config()
            cfg["headless"] = headless

            def on_captcha():
                ev = threading.Event()
                self._msg_queue.put(("captcha", ev))
                ev.wait()

            def on_progress(idx, total, elapsed):
                self._msg_queue.put(("progress", idx, total, elapsed))

            all_data = []
            total_users = len(urls)

            for ui, url in enumerate(urls, 1):
                self._msg_queue.put(("user_progress", ui, total_users))
                print(f"\n══ 用户 {ui}/{total_users} ══")

                scraper = AmazonReviewScraper(profile_url=url, config=cfg)
                try:
                    data = scraper.scrape(
                        progress_callback=on_progress,
                        captcha_callback=on_captcha)
                    AmazonReviewScraper.save_to_json(data)
                    all_data.append(data)
                    print(f"用户 {ui} 完成，{data['total_reviews']} 条评论\n")
                except Exception as e:
                    print(f"用户 {ui} 爬取失败: {e}\n")
                    traceback.print_exc()

            if all_data:
                xlsx_path = AmazonReviewScraper.save_batch_excel(all_data)
                n_users = len(all_data)
                n_reviews = sum(d["total_reviews"] for d in all_data)
                self._msg_queue.put((
                    "done",
                    f"批量爬取完成！\n\n"
                    f"成功: {n_users}/{total_users} 个用户\n"
                    f"总评论: {n_reviews} 条\n"
                    f"Excel: {xlsx_path}"))
            else:
                self._msg_queue.put(("done", "所有用户爬取均失败"))

        except Exception as e:
            print(f"\n爬取失败: {e}")
            traceback.print_exc()
            self._msg_queue.put(("done",))

    # ── 设置 ──

    def _on_settings(self):
        SettingsDialog(self)

    def _on_reset(self):
        if messagebox.askyesno(
                "确认", "确定要重置为默认配置吗？", parent=self):
            save_config(dict(DEFAULT_CONFIG))
            print("已重置为默认配置\n")


# ── 入口点 ────────────────────────────────────────────────────────────────────


def _cli_main():
    while True:
        try:
            choice = show_menu()
            if choice == "1":
                do_login()
            elif choice == "2":
                do_scrape()
            elif choice == "3":
                do_config()
            elif choice == "4":
                do_reset_config()
            elif choice == "0":
                print("再见！")
                break
            else:
                print("无效选择，请重新输入")
        except KeyboardInterrupt:
            print("\n\n已取消操作")
        except Exception as e:
            print(f"\n发生错误: {e}")
            traceback.print_exc()
        print()
        input("按回车返回主菜单...")


def main():
    if "--cli" in sys.argv:
        _cli_main()
    else:
        app = App()
        app.mainloop()


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"\n程序异常: {e}")
        traceback.print_exc()
        input("\n按回车退出...")
